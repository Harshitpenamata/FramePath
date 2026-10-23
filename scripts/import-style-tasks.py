"""Import the style task briefs into app/videography/style-tasks.json.

    python scripts/import-style-tasks.py [docs/sources/Framepath_Style_Layer_Template.xlsx]

Reads "Style task brief" from the "Style layer" tab and the "Should - Working with people" tab. Output is
{module id: {style: task}}; a module or style with no brief keeps the shared task. Example videos and extra
lessons are not imported yet. Needs openpyxl.
"""
import json, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'docs' / 'sources' / 'Framepath_Style_Layer_Template.xlsx'
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
out = {}

def read(tab):
    rows = list(wb[tab].iter_rows(values_only=True))
    cols = {str(h).strip(): i for i, h in enumerate(rows[0]) if h}
    for r in rows[1:]:
        mid, style, brief = (r[cols[k]] for k in ('Module ID', 'Style', 'Style task brief'))
        if mid and style and brief and str(brief).strip():
            out.setdefault(str(mid).strip(), {})[str(style).strip()] = ' '.join(str(brief).split())

read('Style layer')
if 'Should - Working with people' in wb.sheetnames:
    read('Should - Working with people')
dest = ROOT / 'app' / 'videography' / 'style-tasks.json'
dest.write_text(json.dumps(dict(sorted(out.items())), ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
print(f'{sum(len(v) for v in out.values())} style tasks for {len(out)} modules -> {dest.relative_to(ROOT)}')
